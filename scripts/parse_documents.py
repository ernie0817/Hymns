#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import json
import logging
import re
import sqlite3
import sys
from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

try:
    from pypdf import PdfReader
except ImportError as exc:  # pragma: no cover
    raise SystemExit("Missing dependency: pypdf. Install it with: pip install pypdf") from exc

try:
    from docx import Document
except ImportError as exc:  # pragma: no cover
    raise SystemExit("Missing dependency: python-docx. Install it with: pip install python-docx") from exc


BASE_DIR = Path(__file__).resolve().parents[1]
DATA_DIR = BASE_DIR / "data"
PROCESSED_DIR = DATA_DIR / "processed"
JSON_OUTPUT = PROCESSED_DIR / "hymns.json"
DB_OUTPUT = PROCESSED_DIR / "hymns.db"
SUPPORTED_EXTENSIONS = {".docx", ".pdf"}

logger = logging.getLogger("hymn_parser")
logger.setLevel(logging.INFO)


def setup_logging() -> None:
    if logger.handlers:
        return

    formatter = logging.Formatter(
        "%(asctime)s | %(levelname)s | %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )

    stream_handler = logging.StreamHandler(sys.stdout)
    stream_handler.setFormatter(formatter)
    logger.addHandler(stream_handler)
    logger.propagate = False


def normalize_whitespace(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def normalize_newlines(value: str) -> str:
    return value.replace("\r\n", "\n").replace("\r", "\n").strip()


def file_sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as fp:
        for chunk in iter(lambda: fp.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


def discover_documents(data_dir: Path) -> List[Path]:
    if not data_dir.exists():
        logger.warning("Data directory does not exist: %s", data_dir)
        return []

    files: List[Path] = []
    for path in sorted(data_dir.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            if "processed" in path.parts:
                continue
            files.append(path)
    return files


def read_pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))
    pages: List[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text:
            pages.append(text)
    return "\n".join(pages)


def read_docx_text(path: Path) -> str:
    document = Document(str(path))
    paragraphs: List[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            paragraphs.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            merged = " | ".join(cell for cell in cells if cell)
            if merged:
                paragraphs.append(merged)

    return "\n".join(paragraphs)


def infer_category_from_path(path: Path) -> str:
    try:
        rel = path.relative_to(DATA_DIR)
    except ValueError:
        return "未分類"

    parts = rel.parts[:-1]
    if parts:
        first = parts[0]
        if first and first not in {"processed"}:
            return first
    return "未分類"


def extract_id_title_from_filename(path: Path) -> Tuple[str, str]:
    stem = path.stem
    if not stem:
        return "", ""

    title_match = re.search(r"\[(.*?)\]$", stem)
    if title_match:
        title = normalize_whitespace(title_match.group(1))
        suffix = stem[: title_match.start()]
        id_match = re.search(r"(\d{1,5})", suffix)
        if id_match:
            return id_match.group(1), title
        if title:
            return "", title

    id_match = re.search(r"(\d{1,5})\s*(?:\[(.*)\]|$)", stem)
    if id_match:
        hymn_id = id_match.group(1)
        title = normalize_whitespace(id_match.group(2) or stem)
        if title and title != hymn_id:
            return hymn_id, title.replace("_", " ").replace("-", " ")

    digits = re.search(r"(\d{1,5})", stem)
    if digits:
        rest = stem[: digits.start()] + stem[digits.end() :]
        title = normalize_whitespace(rest.replace("_", " ").replace("-", " "))
        if title:
            return digits.group(1), title

    return "", ""


def extract_category(text: str, file_name: str, path: Optional[Path] = None) -> str:
    if path is not None:
        inferred = infer_category_from_path(path)
        if inferred and inferred != "未分類":
            return inferred

    lower_text = text.lower()
    patterns = [
        "大本詩歌",
        "詩歌補充本",
        "補充本",
        "新歌",
        "舊歌",
        "兒童詩歌",
        "敬拜",
        "詩歌",
        "聖歌",
        "台語詩歌",
    ]

    for pattern in patterns:
        if pattern in text or pattern.lower() in lower_text:
            return pattern

    match = re.search(r"(?:分類|類別|主題|category)\s*[:：]?\s*([^\n]+)", text, flags=re.IGNORECASE)
    if match:
        value = normalize_whitespace(match.group(1))
        if value:
            return value

    lower_name = file_name.lower()
    for pattern in patterns:
        if pattern.lower() in lower_name:
            return pattern

    return "未分類"


def extract_id_and_title(text: str, file_stem: str, path: Optional[Path] = None) -> Tuple[str, str]:
    if path is not None:
        hymn_id, title = extract_id_title_from_filename(path)
        if hymn_id or title:
            return hymn_id, title

    lines = [line.strip() for line in normalize_newlines(text).splitlines() if line.strip()]
    for line in lines[:80]:
        if not re.search(r"[\u4e00-\u9fffA-Za-z]", line):
            continue

        match = re.match(r"^(?:[\u4e00-\u9fffA-Za-z\-\s]+)?\s*(\d{1,5})\s*[:：\-–—\s]*\s*(.+)$", line)
        if match:
            hymn_id = match.group(1)
            title = normalize_whitespace(match.group(2))
            if title and not re.fullmatch(r"[\d\s\-：.:/|#♭♮♯]+", title):
                return hymn_id, title

        match = re.match(r"^(\d{1,5})\s*[:：\-–—\s]*\s*(.+)$", line)
        if match:
            hymn_id = match.group(1)
            title = normalize_whitespace(match.group(2))
            if title and not re.fullmatch(r"[\d\s\-：.:/|#♭♮♯]+", title):
                return hymn_id, title

        if re.search(r"[\u4e00-\u9fff]", line):
            title = normalize_whitespace(line)
            if not re.fullmatch(r"[\d\s\-：.:/|#♭♮♯]+", title):
                digits = re.search(r"(\d{1,5})", title)
                if digits:
                    return digits.group(1), title
                return "", title

    digits = re.search(r"(\d{1,5})", file_stem)
    if digits:
        return digits.group(1), normalize_whitespace(file_stem.replace("_", " ").replace("-", " "))

    return "", ""


def extract_lyrics(text: str, hymn_id: str, title: str) -> str:
    lines = text.splitlines()
    cleaned: List[str] = []
    seen_title = False

    for line in lines:
        value = line.strip()
        if not value:
            continue

        meta_pattern = re.compile(
            r"^(?:詩歌\s*編號|歌曲\s*編號|編號|No\.?|NO\.?|编号|ID|分類|類別|主題|標題|歌名|title|category)\s*[:：]?",
            flags=re.IGNORECASE,
        )
        if meta_pattern.match(value):
            continue

        if re.fullmatch(r"\d{1,5}", value):
            continue

        if title and normalize_whitespace(value) == normalize_whitespace(title):
            if not seen_title:
                seen_title = True
                continue

        if hymn_id and value == hymn_id:
            continue

        cleaned.append(line.rstrip())

    lyrics = "\n".join(cleaned).strip()
    if lyrics:
        return normalize_newlines(lyrics)

    # Fallback: keep only substantial content; ignore obvious metadata from all lines
    fallback: List[str] = []
    for line in lines:
        value = line.strip()
        if not value:
            continue
        if re.match(r"^(?:詩歌\s*編號|歌曲\s*編號|編號|No\.?|NO\.?|编号|ID|分類|類別|主題|標題|歌名|title|category)", value, flags=re.IGNORECASE):
            continue
        if re.fullmatch(r"\d{1,5}", value):
            continue
        fallback.append(line.rstrip())

    return normalize_newlines("\n".join(fallback))


def parse_document(path: Path) -> Optional[Dict[str, Any]]:
    try:
        if path.suffix.lower() == ".pdf":
            raw_text = read_pdf_text(path)
        elif path.suffix.lower() == ".docx":
            raw_text = read_docx_text(path)
        else:
            logger.warning("Unsupported file type: %s", path)
            return None

        if not raw_text or not raw_text.strip():
            logger.warning("File is empty or no text extracted: %s", path)
            return None

        text = normalize_newlines(raw_text)
        hymn_id, title = extract_id_and_title(text, path.stem, path)
        if not hymn_id or not title:
            filename_id, filename_title = extract_id_title_from_filename(path)
            if filename_id:
                hymn_id = filename_id
            if filename_title:
                title = filename_title

        category = extract_category(text, path.name, path)
        lyrics = extract_lyrics(text, hymn_id, title)

        if not title:
            title = path.stem or "未命名詩歌"
        if not hymn_id:
            hymn_id = "unknown"

        hymn_id = normalize_whitespace(str(hymn_id))
        title = normalize_whitespace(title)
        category = normalize_whitespace(category)
        lyrics = normalize_newlines(lyrics)

        if not lyrics:
            logger.warning("Lyrics not extracted for %s; skipping file.", path)
            return None

        return {
            "id": hymn_id,
            "title": title,
            "category": category,
            "lyrics": lyrics,
            "source_path": str(path.relative_to(BASE_DIR)),
            "source_file": path.name,
            "hash": file_sha256(path),
            "parsed_at": datetime.now(timezone.utc).isoformat(),
        }
    except Exception as exc:
        logger.exception("Failed to parse document: %s | error=%s", path, exc)
        return None


def create_sqlite_db(db_path: Path) -> sqlite3.Connection:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(db_path))
    conn.execute("PRAGMA journal_mode=WAL;")
    conn.execute("PRAGMA synchronous=NORMAL;")
    conn.execute("PRAGMA foreign_keys=ON;")

    conn.execute("DROP TABLE IF EXISTS hymns_fts;")
    conn.execute("DROP TABLE IF EXISTS hymns;")

    conn.execute(
        """
        CREATE TABLE hymns (
            id TEXT PRIMARY KEY,
            title TEXT NOT NULL,
            category TEXT,
            lyrics TEXT NOT NULL,
            source_path TEXT,
            source_file TEXT,
            hash TEXT,
            parsed_at TEXT
        );
        """
    )

    conn.execute(
        """
        CREATE VIRTUAL TABLE hymns_fts
        USING fts5(
            id UNINDEXED,
            title,
            category,
            lyrics,
            content='hymns',
            content_rowid='rowid'
        );
        """
    )

    conn.execute(
        """
        CREATE TRIGGER hymns_ai AFTER INSERT ON hymns BEGIN
            INSERT INTO hymns_fts(rowid, id, title, category, lyrics)
            VALUES (new.rowid, new.id, new.title, new.category, new.lyrics);
        END;
        """
    )

    conn.execute(
        """
        CREATE TRIGGER hymns_ad AFTER DELETE ON hymns BEGIN
            INSERT INTO hymns_fts(hymns_fts, rowid, id, title, category, lyrics)
            VALUES('delete', old.rowid, old.id, old.title, old.category, old.lyrics);
        END;
        """
    )

    conn.execute(
        """
        CREATE TRIGGER hymns_au AFTER UPDATE ON hymns BEGIN
            INSERT INTO hymns_fts(hymns_fts, rowid, id, title, category, lyrics)
            VALUES('delete', old.rowid, old.id, old.title, old.category, old.lyrics);
            INSERT INTO hymns_fts(rowid, id, title, category, lyrics)
            VALUES (new.rowid, new.id, new.title, new.category, new.lyrics);
        END;
        """
    )

    conn.commit()
    return conn


def populate_database(conn: sqlite3.Connection, hymns: List[Dict[str, Any]]) -> None:
    for hymn in hymns:
        try:
            conn.execute(
                """
                INSERT INTO hymns (
                    id, title, category, lyrics, source_path, source_file, hash, parsed_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    hymn["id"],
                    hymn["title"],
                    hymn["category"],
                    hymn["lyrics"],
                    hymn["source_path"],
                    hymn["source_file"],
                    hymn["hash"],
                    hymn["parsed_at"],
                ),
            )
        except sqlite3.Error as exc:
            logger.exception("Failed to insert hymn into SQLite: %s | error=%s", hymn.get("title"), exc)

    conn.execute("INSERT INTO hymns_fts(hymns_fts) VALUES('rebuild');")
    conn.commit()


def write_json_output(hymns: List[Dict[str, Any]]) -> None:
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
    payload = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "count": len(hymns),
        "items": hymns,
    }

    with JSON_OUTPUT.open("w", encoding="utf-8") as fp:
        json.dump(payload, fp, ensure_ascii=False, indent=2)
        fp.write("\n")


def main() -> int:
    setup_logging()
    logger.info("Starting hymn document parsing. Project root: %s", BASE_DIR)

    try:
        documents = discover_documents(DATA_DIR)
        if not documents:
            logger.warning("No .docx or .pdf files were found in %s", DATA_DIR)
            PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
            write_json_output([])
            logger.info("Created empty JSON output: %s", JSON_OUTPUT)
            return 0

        parsed_hymns: List[Dict[str, Any]] = []
        seen_ids: set[str] = set()

        for document_path in documents:
            logger.info("Parsing file: %s", document_path)
            hymn = parse_document(document_path)
            if hymn is None:
                continue

            hymn_id = str(hymn["id"]).strip()
            if hymn_id and hymn_id in seen_ids:
                logger.warning("Duplicate hymn id detected: %s | file=%s", hymn_id, document_path)
                continue

            if hymn_id:
                seen_ids.add(hymn_id)

            parsed_hymns.append(hymn)

        if not parsed_hymns:
            logger.warning("No valid hymn records were extracted from %s", DATA_DIR)
            PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
            write_json_output([])
            return 0

        write_json_output(parsed_hymns)
        logger.info("JSON output saved to: %s", JSON_OUTPUT)

        conn = create_sqlite_db(DB_OUTPUT)
        try:
            populate_database(conn, parsed_hymns)
        finally:
            conn.close()

        logger.info("SQLite database saved to: %s", DB_OUTPUT)
        logger.info("Finished successfully. Parsed %d hymn records.", len(parsed_hymns))
        return 0

    except Exception as exc:
        logger.exception("Fatal error during parsing process: %s", exc)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
